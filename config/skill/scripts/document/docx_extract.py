"""Extract paragraphs and tables from a DOCX file for workflow use."""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import _textio
import _toolmeta

_textio.force_utf8_stdio()
_TOOL = {
    "id": "docx_extract", "name": "docx_extract", "type": "PYTHON",
    "description": "Extract text and tables from a .docx file. Input JSON: path and optional base_dir.",
    "parameters": {"type": "object", "properties": {
        "purpose": {"type": "string"}, "path": {"type": "string"}, "base_dir": {"type": "string"}},
        "required": ["path"]},
    "metadata": {"domain": "document", "tags": "docx,word,extract"}}
if _toolmeta.describe(_TOOL):
    sys.exit(0)

def main(spec):
    path = spec.get("path")
    if not isinstance(path, str) or not path.strip():
        raise ValueError("path is required")
    base = spec.get("base_dir") or os.environ.get("CLARA_WORK_DIR") or os.getcwd()
    target = path if os.path.isabs(path) else os.path.join(base, path)
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError("python-docx is required: pip install python-docx")
    document = Document(target)
    blocks = [{"type": "paragraph", "text": p.text} for p in document.paragraphs if p.text]
    for index, table in enumerate(document.tables):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        blocks.append({"type": "table", "index": index, "rows": rows})
    return {"path": os.path.abspath(target), "format": "docx", "blocks": blocks}

if __name__ == "__main__":
    try:
        print(json.dumps(main(json.load(sys.stdin)), ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"error": str(exc)}, ensure_ascii=False))
        sys.exit(1)
